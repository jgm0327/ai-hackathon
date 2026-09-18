"""경력기술서 마크다운 → Word(.docx) 변환 (9/16 신규).

Figma "4.2 경력기술서 빌더" 하단의 "Word" 버튼은 그동안 프론트에서 `disabled` +
"준비 중"으로 막혀 있던 스텁이었다 — 실제로 구현한다.

프론트가 이미 화면 렌더링/"마크다운 복사"에 쓰고 있는 것과 동일한 마크다운
문자열(`buildResumeMarkdown()`)을 그대로 받아 변환한다. 백엔드가 STAR 구조를
다시 조합하지 않고 이미 검증된 프론트 텍스트를 그대로 문서화하는 방식이라,
화면에서 보는 내용과 다운로드한 문서 내용이 어긋날 일이 없다.

지원하는 마크다운은 이 앱이 실제로 생성하는 형태로 한정한다 — 범용 마크다운
파서가 아니다(`# 제목`, `## 항목 제목`, `- 불릿`, 나머지는 일반 문단).
"""
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

# 한글 폰트 (9/18 추가). python-docx 기본 템플릿은 `rFonts`를 아예 지정하지 않아서,
# Word가 테마 기본값(Calibri 등 라틴 폰트)으로 렌더링한다 — 한글 글리프가 없어서
# 글자가 깨져 보인다(9/18 실제 발생). 특히 `w:eastAsia`를 따로 지정해야 하는데
# python-docx의 `style.font.name`은 ascii/hAnsi만 설정하므로 XML을 직접 건드린다.
_KOREAN_FONT = "Malgun Gothic"  # 윈도우 기본 한글 폰트. 없는 환경에선 Word가 알아서 대체한다.

# 이 변환기가 실제로 쓰는 스타일만 손본다.
_STYLES_TO_PATCH = ("Normal", "Heading 1", "Heading 2", "List Bullet")


def _apply_korean_font(document: Document) -> None:
    for style_name in _STYLES_TO_PATCH:
        try:
            style = document.styles[style_name]
        except KeyError:
            continue  # 템플릿에 없는 스타일이면 건너뛴다
        style.font.name = _KOREAN_FONT  # ascii / hAnsi
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), _KOREAN_FONT)  # 한글은 이 설정을 따른다

        # 제목 스타일(Heading 1/2)에는 `*Theme` 속성이 함께 들어있는데, OOXML 규격상
        # 테마 속성이 명시적 폰트보다 **우선**한다. 그대로 두면 위에서 지정한 한글
        # 폰트가 무시돼 제목만 계속 깨진다(9/18 실측). 테마 속성을 지워야 적용된다.
        for theme_attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            key = qn(f"w:{theme_attr}")
            if key in rfonts.attrib:
                del rfonts.attrib[key]


def markdown_to_docx_bytes(markdown_text: str) -> bytes:
    document = Document()
    _apply_korean_font(document)
    for raw_line in markdown_text.split("\n"):
        line = raw_line.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
